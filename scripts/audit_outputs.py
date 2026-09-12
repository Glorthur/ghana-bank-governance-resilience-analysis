from __future__ import annotations

import hashlib
import json
import argparse
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "data" / "frozen" / "bank_year_master.csv"
OUTPUT_ROOT = ROOT / "outputs"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def main() -> None:
    global INPUT, OUTPUT_ROOT
    parser = argparse.ArgumentParser(description="Audit analysis outputs.")
    parser.add_argument("--input", type=Path, default=INPUT, help="Input bank-year CSV")
    parser.add_argument("--output", type=Path, default=OUTPUT_ROOT, help="Analysis output directory")
    args = parser.parse_args()
    INPUT = args.input.resolve()
    OUTPUT_ROOT = args.output.resolve()
    required = [
        "input_frozen/bank_year_master.csv",
        "data/analysis_ready_panel.csv",
        "data/primary_model_sample_95.csv",
        "tables/analysis_sample_register.csv",
        "tables/table_01_descriptive_core.csv",
        "tables/table_07_correlation_matrix.csv",
        "tables/table_08_vif.csv",
        "tables/table_09_panel_regressions.csv",
        "tables/table_10_simple_slopes.csv",
        "tables/table_11_robustness_models.csv",
        "tables/table_12_leave_one_bank_out.csv",
        "diagnostics/hausman_test.csv",
        "diagnostics/mundlak_test.csv",
        "diagnostics/wild_cluster_bootstrap_h1_h3.csv",
        "diagnostics/heteroskedasticity.csv",
        "diagnostics/serial_correlation.csv",
        "diagnostics/cross_sectional_dependence.csv",
        "figures/figure_01_ori_distribution.png",
        "figures/figure_02_ori_year_trend.png",
        "figures/figure_03_marginal_effect_board_independence.png",
        "figures/figure_04_leave_one_bank_out_interaction.png",
        "ANALYSIS_TABLES.xlsx",
        "ANALYSIS_INTERPRETATION_MEMO.md",
        "CHAPTER_FOUR_DRAFTING_OUTLINE.md",
        "DATA_VALIDATION_REPORT.md",
        "manifests/VALIDATION_REPORT.json",
        "manifests/OUTPUT_MANIFEST.json",
    ]
    checks: dict[str, object] = {}
    checks["all_required_files_present"] = all(
        (OUTPUT_ROOT / relative).is_file() and (OUTPUT_ROOT / relative).stat().st_size > 0
        for relative in required
    )
    missing = [relative for relative in required if not (OUTPUT_ROOT / relative).is_file()]

    frozen = OUTPUT_ROOT / "input_frozen" / "bank_year_master.csv"
    checks["frozen_hash_matches_input"] = frozen.is_file() and INPUT.is_file() and sha256(frozen) == sha256(INPUT)

    panel = pd.read_csv(OUTPUT_ROOT / "data" / "analysis_ready_panel.csv")
    primary = pd.read_csv(OUTPUT_ROOT / "data" / "primary_model_sample_95.csv")
    register = pd.read_csv(OUTPUT_ROOT / "tables" / "analysis_sample_register.csv")
    checks["panel_rows_96"] = len(panel) == 96
    checks["panel_unique_keys"] = not panel.duplicated(["bank_id", "year"]).any()
    checks["register_rows_96"] = len(register) == 96
    checks["primary_rows_95"] = len(primary) == 95
    checks["primary_banks_16"] = primary["bank_id"].nunique() == 16
    source_ceiling = int((pd.read_csv(frozen)["ori_score"] == 100).sum())
    checks["ori_ceiling_matches_frozen_source"] = (
        int((primary["ori_score"] == 100).sum()) == source_ceiling
    )

    regressions = pd.read_csv(OUTPUT_ROOT / "tables" / "table_09_panel_regressions.csv")
    expected_models = {
        "pooled_main",
        "fe_main",
        "re_main",
        "pooled_moderation",
        "fe_moderation",
        "re_moderation",
    }
    checks["six_primary_models"] = set(regressions["model"].unique()) == expected_models
    checks["all_primary_models_n95"] = (regressions["observations"] == 95).all()
    checks["all_primary_models_year_effects"] = regressions["year_effects"].astype(bool).all()
    fe_hypotheses = regressions.loc[
        (regressions["model"] == "fe_moderation")
        & regressions["variable"].isin(
            ["board_independence_c", "firm_size_c", "board_x_size"]
        )
    ]
    checks["fe_has_h1_h3"] = len(fe_hypotheses) == 3
    checks["fe_h1_h3_finite"] = np.isfinite(
        fe_hypotheses[["coefficient", "std_error", "p_value"]].to_numpy()
    ).all()

    wild = pd.read_csv(OUTPUT_ROOT / "diagnostics" / "wild_cluster_bootstrap_h1_h3.csv")
    checks["wild_tests_h1_h3"] = set(wild["hypothesis"]) == {"H1", "H2", "H3"}
    checks["wild_tests_estimated"] = (wild["status"] == "estimated").all()
    checks["wild_pvalues_finite"] = np.isfinite(wild["p_value"]).all()
    checks["wild_statistics_finite"] = np.isfinite(wild["t_statistic"]).all()

    robustness = pd.read_csv(OUTPUT_ROOT / "tables" / "table_11_robustness_models.csv")
    checks["car_n70_b15"] = (
        robustness.loc[robustness["model"] == "CAR-augmented two-way FE", "observations"].eq(70).all()
        and robustness.loc[robustness["model"] == "CAR-augmented two-way FE", "banks"].eq(15).all()
    )
    checks["bounded_models_n95_b16"] = (
        robustness.loc[
            robustness["model"].isin(
                ["Full-disclosure logit GLM", "Fractional-response logit GLM"]
            ),
            "observations",
        ].eq(95).all()
        and robustness.loc[
            robustness["model"].isin(
                ["Full-disclosure logit GLM", "Fractional-response logit GLM"]
            ),
            "banks",
        ].eq(16).all()
    )
    checks["npl_n43_b10"] = (
        robustness.loc[
            robustness["model"] == "NPL separate-outcome two-way FE", "observations"
        ].eq(43).all()
        and robustness.loc[
            robustness["model"] == "NPL separate-outcome two-way FE", "banks"
        ].eq(10).all()
    )

    leave = pd.read_csv(OUTPUT_ROOT / "tables" / "table_12_leave_one_bank_out.csv")
    checks["leave_one_out_16"] = len(leave) == 16 and (leave["status"] == "estimated").all()

    figure_checks = {}
    for path in sorted((OUTPUT_ROOT / "figures").glob("*.png")):
        with Image.open(path) as image:
            pixels = np.asarray(image.convert("RGB"), dtype=float)
            figure_checks[path.name] = {
                "width": image.width,
                "height": image.height,
                "pixel_std": float(pixels.std()),
                "nonblank": bool(image.width >= 800 and image.height >= 500 and pixels.std() > 5),
            }
    checks["all_figures_nonblank"] = len(figure_checks) == 4 and all(
        value["nonblank"] for value in figure_checks.values()
    )

    workbook = load_workbook(OUTPUT_ROOT / "ANALYSIS_TABLES.xlsx", read_only=True, data_only=True)
    expected_sheets = {
        "Descriptives",
        "By Year",
        "ORI Components",
        "Correlations",
        "VIF",
        "Panel Models",
        "Simple Slopes",
        "Robustness",
        "Leave One Out",
        "Hausman",
        "Mundlak",
        "Wild Bootstrap",
    }
    checks["workbook_opens_with_expected_sheets"] = expected_sheets.issubset(
        set(workbook.sheetnames)
    )
    workbook.close()

    manifest = json.loads(
        (OUTPUT_ROOT / "manifests" / "OUTPUT_MANIFEST.json").read_text(encoding="utf-8")
    )
    manifest_checks = []
    for item in manifest["files"]:
        normalized_path = item["path"].replace("\\", "/")
        if normalized_path.startswith("second_pass_validation/") or normalized_path == "logs/run_summary.json":
            continue
        path = OUTPUT_ROOT / item["path"]
        manifest_checks.append(path.is_file() and sha256(path) == item["sha256"])
    checks["manifest_hashes_match"] = all(manifest_checks)

    bool_checks = [bool(value) for value in checks.values() if isinstance(value, (bool, np.bool_))]
    result = "PASS" if all(bool_checks) else "FAIL"
    payload = {
        "result": result,
        "checks": checks,
        "missing_required_files": missing,
        "figure_checks": figure_checks,
    }
    manifest_dir = OUTPUT_ROOT / "manifests"
    manifest_dir.mkdir(parents=True, exist_ok=True)
    (manifest_dir / "FINAL_ANALYSIS_AUDIT.json").write_text(
        json.dumps(
            payload,
            indent=2,
            default=lambda value: value.item() if isinstance(value, np.generic) else str(value),
        ),
        encoding="utf-8",
    )
    print(
        json.dumps(
            payload,
            indent=2,
            default=lambda value: value.item() if isinstance(value, np.generic) else str(value),
        )
    )
    if result != "PASS":
        raise SystemExit(1)


if __name__ == "__main__":
    main()
